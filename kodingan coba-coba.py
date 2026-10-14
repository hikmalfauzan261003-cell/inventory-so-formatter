import csv
import io
import posixpath
import re
import zipfile
from copy import copy

import openpyxl
import pandas as pd
import requests
import streamlit as st
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

st.set_page_config(
    page_title="Stock Opname Khusus SUki", page_icon="member.jpeg", layout="wide"
)
col4, col5, col6 = st.columns ([1, 2, 1])
with col4:
    st.image("member.jpeg", width = 200)
with col5:
    st.title("Stock Opname Worksuki Generator")
with col6:
    st.image("member.jpeg", width = 200)
st.write(
    "Halo Para Suki dan Member, Anda malas bikin worksheet? Sama saya juga makanya saya bikin ini."
)

st.divider()

# ---------------------------------------------------------
# DIRECT LINK TEMPLATE MASTER GOOGLE DRIVE (ANTI VIRUS WARNING)
# ---------------------------------------------------------
FILE_ID = "13HVdBlxbyBhpwTLxNwfFV5NXapLf0Tli"


@st.cache_data
def fetch_master_template():
    """Mengunduh template master dari Google Drive dengan penanganan konfirmasi virus scan."""
    URL = "https://docs.google.com/uc?export=download"
    session = requests.Session()
    response = session.get(URL, params={"id": FILE_ID}, stream=True)
    
    # Cek token konfirmasi jika file besar / kena virus warning Google Drive
    for key, value in response.cookies.items():
        if key.startswith("download_warning"):
            response = session.get(URL, params={"id": FILE_ID, "confirm": value}, stream=True)
            break
            
    response.raise_for_status()
    return io.BytesIO(response.content)

# ---------------------------------------------------------
# STEP 1: UPLOAD DOKUMEN
# ---------------------------------------------------------
st.header("1. Aplot Doksli Dulu Lohya 😹")

col1, col2, col3 = st.columns(3)

with col1:
    csv_file = st.file_uploader(
        "Upload Doksli Batmis 😹 (.csv / .xlsx)",
        type=["csv", "xlsx", "xls"],
        help="Upload file dari batmis (bisa format CSV atau Excel).",
    )

with col2:
    prev_so_file = st.file_uploader(
        "Upload Dokumen Prev SO 🟣🟣🟣 (.xlsx)",
        type=["xlsx"],
        help="Note : Worksheet utama harus direname jadi 'Worksheet' (W besar sisanya kecil)",
    )

with col3:
    missing_file = st.file_uploader(
        "Upload Doksli Missing (.xlsx / .csv)",
        type=["xlsx", "xls", "csv"],
        help="Note : Delete kolom paling atas & kiri, delete kolom O sampe ujung kanan.",
    )

st.divider()

# ---------------------------------------------------------
# STEP 2: INPUT METADATA & SUMMARY
# ---------------------------------------------------------
st.header("2. Input Informasi Metadata & Summary")

col_m1, col_m2, col_m3 = st.columns(3)
with col_m1:
    station_input = st.text_input("Station", value="SUB")
with col_m2:
    location_input = st.text_input(
        "Location", value="S1, K86, K88, K87, K56 & K58"
    )
with col_m3:
    periode_input = st.text_input("Periode Date", value="SEPTEMBER 2026")

st.subheader("Pengaturan Sheet Summary (Dinamis Per Lokasi)")
st.caption("Masukkan detail divisi, PIC, kode lokasi, dan deskripsi lokasi.")

if "summary_rows" not in st.session_state:
    st.session_state.summary_rows = [
        {
            "division": "LINE MAINTENANCE",
            "pic": "Bahlil",
            "loc_code": "K86",
            "loc_desc": "STORE SUB",
        }
    ]

# Tombol Tambah / Hapus Lokasi
col_b1, col_b2, _ = st.columns([1, 1, 4])
with col_b1:
    if st.button("➕ Tambah Lokasi"):
        st.session_state.summary_rows.append(
            {"division": "", "pic": "", "loc_code": "", "loc_desc": ""}
        )
        st.rerun()

with col_b2:
    if (
        st.button("➖ Hapus Lokasi Terakhir")
        and len(st.session_state.summary_rows) > 1
    ):
        st.session_state.summary_rows.pop()
        st.rerun()

# Form dinamis untuk input summary
summary_data_inputs = []
for idx, row in enumerate(st.session_state.summary_rows):
    st.markdown(f"**Lokasi #{idx + 1}**")
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        div = st.text_input(
            f"Division #{idx+1}", value=row["division"], key=f"div_{idx}"
        )
    with c2:
        pic = st.text_input(
            f"PIC #{idx+1}", value=row["pic"], key=f"pic_{idx}"
        )
    with c3:
        loc = st.text_input(
            f"LOC Code #{idx+1}", value=row["loc_code"], key=f"loc_{idx}"
        )
    with c4:
        desc = st.text_input(
            f"LOC Description #{idx+1}",
            value=row["loc_desc"],
            key=f"desc_{idx}",
        )
    summary_data_inputs.append(
        {"division": div, "pic": pic, "loc_code": loc, "loc_desc": desc}
    )

st.divider()

# ---------------------------------------------------------
# STEP 3: HELPER FUNCTIONS & EXEKUSI
# ---------------------------------------------------------
st.header("3. Eksekusi & Pemrosesan")

# ---------------------------------------------------------
# SAKLAR FITUR TAMBAHAN (ubah True / False kalau perlu diuji)
# ---------------------------------------------------------
# True  = logo header cetak dikembalikan ke file hasil (openpyxl biasanya membuangnya)
RESTORE_HEADER_LOGO = True
# True  = PivotTable di-refresh otomatis saat file dibuka di Excel
#         (False = pivot tetap data lama sampai kamu klik Data > Refresh All)
REFRESH_PIVOT_ON_OPEN = False

# Karakter kontrol yang ilegal di XML (penyebab file Excel corrupt)
ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# Posisi kolom bawaan (sesuai TEMPLATE_SO.xlsx). Dipakai hanya sebagai cadangan
# kalau header di baris 7 tidak ditemukan. Prioritas utama: baca dari header.
HEADER_ROW = 7
FIRST_DATA_ROW = 8
DEFAULT_COLS = {
    "no": 1, "count_no": 2, "loc": 3, "bin": 4, "grb": 5, "batch": 6,
    "pn": 7, "sn": 8, "pn_desc": 9, "qty_avail": 10, "qty_resv": 11,
    "qty_trans": 12, "qty_pending": 13, "qty_us": 14, "qty_repair": 15,
    "shelf": 16, "condition": 17, "category": 18, "owner": 19, "uom": 20,
    "emro": 21, "actual": 22, "diff": 23, "result": 24, "status": 25,
    "prev_so": 30, "prev_status": 31, "cat": 32,
}
# Teks header (huruf besar, spasi dirapikan) -> nama field
HEADER_TO_FIELD = {
    "NO": "no", "COUNT NO": "count_no", "LOC": "loc", "BIN": "bin",
    "GRB": "grb", "BATCH": "batch", "PN": "pn", "SN": "sn",
    "PN DESCRIPTION": "pn_desc", "QTY AVAILABLE": "qty_avail",
    "QTY RESERVED": "qty_resv", "QTY IN TRANSFER": "qty_trans",
    "QTY PENDING R/I": "qty_pending", "QTY US": "qty_us",
    "QTY IN REPAIR": "qty_repair", "SHELF LIFE EXP": "shelf",
    "CONDITION": "condition", "CATEGORY": "category", "OWNER": "owner",
    "UOM": "uom", "QTY EMRO": "emro", "QTY ACTUAL": "actual",
    "DIFF": "diff", "RESULT": "result", "STATUS": "status",
    "PREV SO": "prev_so", "PREV STATUS": "prev_status", "CAT": "cat",
}


def normalize_batch(val):
    """
    Fungsi Normalisasi Agresif & Kebal Tipe Data:
    - Mengonversi integer, float, string menjadi String murni.
    - Menghapus .0 di akhir angka float (misal: 4559080.0 -> 4559080).
    - Menghapus SEMUA jenis spasi, newline (\n), tab (\t), dan karakter invisible (\xa0).
    - Mengubah semua huruf menjadi UPPERCASE.
    """
    if pd.isna(val) or val is None:
        return ""

    val_str = str(val).strip()

    # Potong desimal .0 di akhir jika ada
    if val_str.endswith(".0"):
        val_str = val_str[:-2]

    # Hapus seluruh whitespace & karakter khusus tersembunyi
    val_str = re.sub(r"\s+", "", val_str).upper()

    return val_str


def load_raw_inventory_data(uploaded_file_obj):
    """Membaca file data mentah inventory, menangani CSV tersembunyi di Excel, & melakukan sorting A-Z."""
    fname = uploaded_file_obj.name.lower()
    df = None

    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        try:
            df = pd.read_excel(uploaded_file_obj)
            if df.shape[1] <= 2:
                uploaded_file_obj.seek(0)
                df_temp = pd.read_excel(uploaded_file_obj)
                non_null_series = df_temp.dropna(how="all").iloc[:, 0].astype(
                    str
                )
                text_data = "\n".join(non_null_series)
                df = pd.read_csv(
                    io.StringIO(text_data),
                    sep=";",
                    quoting=csv.QUOTE_MINIMAL,
                    on_bad_lines="skip",
                )
        except Exception:
            uploaded_file_obj.seek(0)
            df = pd.read_excel(uploaded_file_obj)
    else:
        try:
            df = pd.read_csv(uploaded_file_obj, sep=";", on_bad_lines="skip")
            if df.shape[1] <= 1:
                uploaded_file_obj.seek(0)
                df = pd.read_csv(
                    uploaded_file_obj, sep=",", on_bad_lines="skip"
                )
        except Exception:
            uploaded_file_obj.seek(0)
            df = pd.read_csv(uploaded_file_obj, sep=",", on_bad_lines="skip")

    df.columns = df.columns.astype(str).str.strip().str.upper()

    # --- FITUR SORTING A-Z (LOC -> BIN -> PN -> SN) ---
    sort_cols = []
    for col in [
        "LOCATION",
        "LOC",
        "BIN",
        "PN",
        "PART NO",
        "PART_NO",
        "SN",
        "SERIAL NO",
        "SERIAL_NO",
    ]:
        if col in df.columns and col not in sort_cols:
            sort_cols.append(col)

    if sort_cols:
        df = df.sort_values(by=sort_cols, ascending=True).reset_index(
            drop=True
        )

    return df


def extract_missing_batch_set(missing_file_obj, target_loc_list=None):
    """
    Mengekstrak Batch MISSING yang HANYA SESUAI dengan kode lokasi
    yang terdaftar pada Form Dinamis Sheet Summary.
    """
    if not missing_file_obj:
        return set()

    missing_batches = set()
    fname = missing_file_obj.name.lower()

    try:
        if fname.endswith(".csv"):
            df_missing = pd.read_csv(
                missing_file_obj,
                sep=None,
                engine="python",
                on_bad_lines="skip",
            )
            df_list = [df_missing]
        else:
            xls = pd.ExcelFile(missing_file_obj)
            df_list = [
                pd.read_excel(xls, sheet_name=s) for s in xls.sheet_names
            ]

        valid_locations = []
        if target_loc_list:
            valid_locations = [
                str(loc).strip().upper()
                for loc in target_loc_list
                if str(loc).strip() != ""
            ]

        for df_sheet in df_list:
            if df_sheet.empty:
                continue

            df_sheet.columns = [
                str(c).strip().upper() for c in df_sheet.columns
            ]

            # Filter lokasi spesifik
            if valid_locations and "LOCATION" in df_sheet.columns:
                loc_pattern = "|".join(
                    [rf"\b{re.escape(loc)}\b" for loc in valid_locations]
                )
                df_sheet = df_sheet[
                    df_sheet["LOCATION"]
                    .astype(str)
                    .str.strip()
                    .str.upper()
                    .str.contains(loc_pattern, regex=True, na=False)
                ]

            batch_cols = [
                c
                for c in df_sheet.columns
                if any(
                    k == c or f" {k}" in f" {c}"
                    for k in ["BATCH", "GRB", "BATCH NO", "BATCH_NO"]
                )
            ]

            if not batch_cols and len(df_sheet.columns) > 0:
                batch_cols = [df_sheet.columns[0]]

            for col in batch_cols:
                for val in df_sheet[col].dropna():
                    norm_v = normalize_batch(val)
                    if norm_v != "":
                        missing_batches.add(norm_v)

    except Exception as e:
        st.warning(f"Catatan pemrosesan file missing: {e}")

    return missing_batches


def get_val(row, *possible_keys, default=None):
    """Mencari nilai kolom secara case-insensitive & aman dari NaN."""
    for key in possible_keys:
        k_upper = key.strip().upper()
        if k_upper in row and pd.notna(row[k_upper]):
            val = row[k_upper]
            return val
    return default


def clean_value(val):
    """
    Membersihkan nilai sebelum ditulis ke Excel supaya file tidak corrupt:
    - NaN / NaT / None / string kosong / string spasi saja -> None (sel benar-benar kosong)
    - Buang karakter kontrol ilegal XML
    - Tipe numpy -> tipe Python biasa, float bulat (4722404.0) -> int
    - Timestamp pandas -> datetime Python
    """
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(val, str):
        val = ILLEGAL_XML_CHARS.sub("", val)
        if val.strip() == "":
            return None
        return val

    if isinstance(val, pd.Timestamp):
        return val.to_pydatetime()

    if hasattr(val, "item"):  # numpy scalar
        try:
            val = val.item()
        except Exception:
            pass

    if isinstance(val, float) and val.is_integer():
        return int(val)
    return val


def set_cell_safe(ws, row, col, value, formula=False):
    """Mengisi sel tanpa merusak jika menimpa MergedCell. Nilai kosong ditulis sebagai sel kosong."""
    value = clean_value(value)
    cell = ws.cell(row=row, column=col)
    if type(cell).__name__ == "MergedCell":
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                ws.cell(
                    row=merged_range.min_row, column=merged_range.min_col
                ).value = value
                break
    else:
        cell.value = value
        # Teks data yang diawali "=" jangan dibaca sebagai rumus
        if not formula and isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"


def build_prev_so_dict(prev_so_file_obj):
    """Pindai secara otomatis baris header & ekstraksi VLOOKUP Prev SO & Status secara presisi."""
    try:
        df_raw = pd.read_excel(
            prev_so_file_obj, sheet_name="Worksheet", header=None
        )
    except Exception:
        prev_so_file_obj.seek(0)
        df_raw = pd.read_excel(prev_so_file_obj, header=None)

    header_idx = None
    for idx, r in df_raw.iterrows():
        row_str_vals = [str(v).strip().upper() for v in r.values if pd.notna(v)]
        if "BATCH" in row_str_vals and (
            "RESULT" in row_str_vals or "STATUS" in row_str_vals
        ):
            header_idx = idx
            break

    prev_dict = {}
    if header_idx is not None:
        prev_so_file_obj.seek(0)
        try:
            df_prev = pd.read_excel(
                prev_so_file_obj, sheet_name="Worksheet", header=header_idx
            )
        except Exception:
            prev_so_file_obj.seek(0)
            df_prev = pd.read_excel(prev_so_file_obj, header=header_idx)

        df_prev.columns = [str(c).strip().upper() for c in df_prev.columns]

        for _, row in df_prev.iterrows():
            batch_val = row.get("BATCH")
            norm_batch = normalize_batch(batch_val)
            if norm_batch != "":
                res_val = (
                    row.get("RESULT")
                    if pd.notna(row.get("RESULT"))
                    else None
                )
                stat_val = (
                    row.get("STATUS")
                    if pd.notna(row.get("STATUS"))
                    else None
                )

                prev_dict[norm_batch] = {
                    "prev_so": res_val,
                    "prev_status": stat_val,
                }
    return prev_dict


def build_column_map(ws):
    """
    Membaca header baris 7 di template untuk menentukan posisi tiap kolom.
    Jadi kalau urutan kolom di template berubah, kode ikut menyesuaikan.
    Kalau sebuah header tidak ketemu, dipakai posisi bawaan (DEFAULT_COLS).
    """
    cols = dict(DEFAULT_COLS)
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(row=HEADER_ROW, column=c).value
        if raw is None:
            continue
        header = re.sub(r"\s+", " ", str(raw)).strip().upper()
        field = HEADER_TO_FIELD.get(header)
        if field:
            cols[field] = c
    return cols


def get_main_table(ws):
    """
    Mengambil OBJEK Excel Table di sheet Worksheet (Table3, atau Table pertama).
    Sengaja memakai dict.values() karena TableList milik openpyxl meng-override
    items() sehingga mengembalikan teks (nama, ref), bukan objek Table.
    """
    try:
        candidates = list(dict.values(ws.tables))
    except Exception:
        candidates = []
    tables = [
        t for t in candidates if not isinstance(t, str) and hasattr(t, "ref")
    ]
    if not tables:
        return None
    for t in tables:
        if getattr(t, "name", None) == "Table3":
            return t
    return tables[0]


def write_so_table_to_sheet(ws, df_data, prev_so_map):
    """Mengisi data ke sheet dari baris 8 ke bawah, lalu menyesuaikan ukuran Table."""
    cols = build_column_map(ws)
    table = get_main_table(ws)

    if table is not None:
        min_c, _min_r, max_c, _max_r = range_boundaries(table.ref)
    else:
        min_c, max_c = 1, max(cols.values())

    # Simpan gaya (border, format angka, dll) baris contoh (baris 8) untuk dipakai ulang
    row_styles = {
        c: copy(ws.cell(row=FIRST_DATA_ROW, column=c)._style)
        for c in range(min_c, max_c + 1)
    }

    # Kosongkan isi baris 8 ke bawah TANPA delete_rows (delete_rows merusak Table/merge)
    last_existing = max(ws.max_row, FIRST_DATA_ROW)
    for r in range(FIRST_DATA_ROW, last_existing + 1):
        for c in range(1, max(ws.max_column, max_c) + 1):
            cell = ws.cell(row=r, column=c)
            if type(cell).__name__ != "MergedCell":
                cell.value = None

    L = lambda key: get_column_letter(cols[key])  # noqa: E731

    for i, (_, row_data) in enumerate(df_data.iterrows()):
        row_idx = FIRST_DATA_ROW + i

        # Terapkan gaya baris contoh ke seluruh kolom Table
        for c in range(min_c, max_c + 1):
            ws.cell(row=row_idx, column=c)._style = copy(row_styles[c])

        raw_batch = get_val(
            row_data, "BATCH", "BATCH NO", "BATCH_NO", default=""
        )
        batch_num = normalize_batch(raw_batch)

        def put(field, value, formula=False):
            set_cell_safe(ws, row_idx, cols[field], value, formula=formula)

        put("no", i + 1)
        put("count_no", get_val(row_data, "COUNT NO", "COUNT_NO"))
        put("loc", get_val(row_data, "LOCATION", "LOC"))
        put("bin", get_val(row_data, "BIN"))
        put("grb", get_val(row_data, "GRB", "GRB NO", "GRB_NO"))
        put("batch", batch_num)
        put("pn", get_val(row_data, "PN", "PART NO", "PART_NO"))
        put("sn", get_val(row_data, "SN", "SERIAL NO", "SERIAL_NO"))
        put(
            "pn_desc",
            get_val(row_data, "PN DESCRIPTION", "DESCRIPTION", "PN DESC"),
        )

        # QTY Columns
        put(
            "qty_avail",
            get_val(row_data, "QTY AVAILABLE", "QTY_AVAIL", "QTY AVAIL", default=0),
        )
        put(
            "qty_resv",
            get_val(row_data, "QTY RESERVED", "QTY_RESV", "QTY RESV", default=0),
        )
        put(
            "qty_trans",
            get_val(row_data, "QTY IN TRANSFER", "QTY_TRANS", "QTY TRANS", default=0),
        )
        put(
            "qty_pending",
            get_val(
                row_data,
                "QTY PENDING RI",
                "QTY PENDING R/I",
                "QTY_PENDING",
                default=0,
            ),
        )
        put("qty_us", get_val(row_data, "QTY US", "QTY_US", default=0))
        put(
            "qty_repair",
            get_val(row_data, "QTY IN REPAIR", "QTY_REPAIR", "QTY REPAIR", default=0),
        )

        put(
            "shelf",
            get_val(
                row_data,
                "SHELF LIFE EXPIRATION",
                "SHELF LIFE EXP",
                "TOOL LIFE EXPIRATION",
            ),
        )
        put("condition", get_val(row_data, "CONDITION", default="SV"))
        put("category", get_val(row_data, "CATEGORY"))
        put("owner", get_val(row_data, "OWNER"))
        put("uom", get_val(row_data, "UOM", default="EA"))

        # --- FORMULA EXCEL DINAMIS (huruf kolom mengikuti header template) ---
        emro, actual = L("emro"), L("actual")
        result, diff = L("result"), L("diff")
        put(
            "emro",
            f"={L('qty_avail')}{row_idx}+{L('qty_resv')}{row_idx}"
            f"+{L('qty_pending')}{row_idx}+{L('qty_us')}{row_idx}",
            formula=True,
        )  # Qty eMRO
        put("actual", None)  # Qty Actual (BLANK)
        put("diff", f"={actual}{row_idx}-{emro}{row_idx}", formula=True)
        put(
            "result",
            f'=IF({emro}{row_idx}=0,"BUG EMRO??/MISSING??",'
            f'IF({actual}{row_idx}=0,"NOT FOUND",'
            f'IF({actual}{row_idx}>{emro}{row_idx},"SURPLUS",'
            f'IF({actual}{row_idx}<{emro}{row_idx},"MINUS","MATCHED"))))',
            formula=True,
        )
        put(
            "status",
            f'=IF({result}{row_idx}="MATCHED","MATCHED","OPEN")',
            formula=True,
        )

        # Kolom manual (Date, Auditor, Remark, Penyelesaian, Corrective Action, dst)
        # dibiarkan kosong. Tidak perlu ditulis sama sekali.

        # --- PREV SO, PREV STATUS, CAT ---
        prev_info = prev_so_map.get(
            batch_num, {"prev_so": None, "prev_status": None}
        )
        put("prev_so", prev_info["prev_so"])
        put("prev_status", prev_info["prev_status"])
        put("cat", get_val(row_data, "CAT", "CATEGORY"))

    # --- Sesuaikan ukuran Table dengan jumlah data (supaya rumus/pivot/drag ikut) ---
    if table is not None:
        last_row = FIRST_DATA_ROW + max(len(df_data), 1) - 1
        new_ref = f"{get_column_letter(min_c)}{HEADER_ROW}:{get_column_letter(max_c)}{last_row}"
        try:
            table.ref = new_ref
            auto_filter = getattr(table, "autoFilter", None)
            if auto_filter is not None and not isinstance(auto_filter, str):
                auto_filter.ref = new_ref
        except Exception as e:
            st.warning(
                f"Ukuran Table tidak bisa disesuaikan otomatis ({e}). "
                "Data tetap terisi, tapi Table mungkin perlu di-resize manual di Excel."
            )


def refresh_pivots_on_open(wb):
    """Minta Excel me-refresh PivotTable saat file dibuka (cache lama di template tidak cocok dengan data baru)."""
    for ws_p in wb.worksheets:
        for pivot in getattr(ws_p, "_pivots", []):
            try:
                pivot.cache.refreshOnLoad = True
            except Exception:
                pass


def restore_header_logo(template_bytes, output_bytes):
    """
    openpyxl membuang gambar (logo) di header cetak. Fungsi ini mengembalikannya
    dengan menyalin bagian header-drawing dari template ke file hasil.
    Kalau ada yang gagal / hasilnya tidak valid, file hasil asli dikembalikan apa adanya.
    Mengembalikan (bytes_file, status). status = "OK" kalau logo berhasil dipulihkan,
    selain itu berisi alasan kenapa tidak.
    """
    NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
    VML_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing"
    try:
        from xml.etree import ElementTree as ET

        zt = zipfile.ZipFile(io.BytesIO(template_bytes))
        zo = zipfile.ZipFile(io.BytesIO(output_bytes))

        def resolve(base_dir, target):
            if target.startswith("/"):
                return target.lstrip("/")
            return posixpath.normpath(posixpath.join(base_dir, target))

        def sheet_paths(z):
            wbx = ET.fromstring(z.read("xl/workbook.xml"))
            rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            rid_target = {r.get("Id"): r.get("Target") for r in rels}
            out = {}
            for s in wbx.find(f"{{{NS_MAIN}}}sheets"):
                tgt = rid_target[s.get(f"{{{NS_REL}}}id")]
                out[s.get("name")] = resolve("xl", tgt)
            return out

        def rels_path(part_path):
            d, f = posixpath.split(part_path)
            return posixpath.join(d, "_rels", f + ".rels")

        t_sheets, o_sheets = sheet_paths(zt), sheet_paths(zo)
        names_in_zip = set(zo.namelist())
        overrides, additions = {}, {}
        content_types = zo.read("[Content_Types].xml").decode("utf-8")
        n = 0

        for sheet_name, t_path in t_sheets.items():
            t_xml = zt.read(t_path).decode("utf-8")
            m = re.search(r'<legacyDrawingHF\s+r:id="([^"]+)"', t_xml)
            if not m or sheet_name not in o_sheets:
                continue
            o_path = o_sheets[sheet_name]
            o_xml = overrides.get(o_path, zo.read(o_path).decode("utf-8"))
            if "legacyDrawingHF" in o_xml:
                continue

            # Cari file VML header di template
            t_rels = ET.fromstring(zt.read(rels_path(t_path)))
            vml_target = next(
                r.get("Target") for r in t_rels if r.get("Id") == m.group(1)
            )
            vml_path = resolve(posixpath.dirname(t_path), vml_target)
            vml_rels_path = rels_path(vml_path)

            n += 1
            new_vml = f"xl/drawings/hfDrawing{n}.vml"
            additions[new_vml] = zt.read(vml_path)

            # Salin gambar yang dipakai VML + buat rels baru
            if vml_rels_path in zt.namelist():
                v_rels = ET.fromstring(zt.read(vml_rels_path))
                new_rel_xml = [f'<Relationships xmlns="{PKG_REL}">']
                for r in v_rels:
                    img_src = resolve(posixpath.dirname(vml_path), r.get("Target"))
                    ext = posixpath.splitext(img_src)[1].lower().lstrip(".")
                    img_new = f"xl/media/hf_logo{n}_{posixpath.basename(img_src)}"
                    if img_new in names_in_zip:
                        raise ValueError("nama gambar bentrok")
                    additions[img_new] = zt.read(img_src)
                    new_rel_xml.append(
                        f'<Relationship Id="{r.get("Id")}" Type="{r.get("Type")}" '
                        f'Target="/{img_new}"/>'
                    )
                    mime = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif"}.get(ext)
                    if mime and f'Extension="{ext}"' not in content_types:
                        content_types = content_types.replace(
                            "</Types>",
                            f'<Default Extension="{ext}" ContentType="{mime}"/></Types>',
                        )
                new_rel_xml.append("</Relationships>")
                additions[rels_path(new_vml)] = "".join(new_rel_xml).encode("utf-8")

            if 'Extension="vml"' not in content_types:
                content_types = content_types.replace(
                    "</Types>",
                    '<Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/></Types>',
                )

            # Hubungkan VML ke sheet hasil
            o_rels_path = rels_path(o_path)
            rid = f"rIdHF{n}"
            if o_rels_path in names_in_zip or o_rels_path in overrides:
                o_rels_xml = overrides.get(o_rels_path) or zo.read(o_rels_path).decode("utf-8")
            else:
                o_rels_xml = f'<Relationships xmlns="{PKG_REL}"></Relationships>'
            o_rels_xml = o_rels_xml.replace(
                "</Relationships>",
                f'<Relationship Id="{rid}" Type="{VML_TYPE}" Target="/{new_vml}"/></Relationships>',
            )
            overrides[o_rels_path] = o_rels_xml

            # Sisipkan <legacyDrawingHF> sebelum tableParts / extLst / akhir worksheet
            tag = f'<legacyDrawingHF xmlns:r="{NS_REL}" r:id="{rid}"/>'
            for marker in ("<tableParts", "<extLst", "</worksheet>"):
                pos = o_xml.find(marker)
                if pos != -1:
                    o_xml = o_xml[:pos] + tag + o_xml[pos:]
                    break
            overrides[o_path] = o_xml

        if n == 0:
            return output_bytes, "Template tidak punya logo header (cek file template di Google Drive)"

        overrides["[Content_Types].xml"] = content_types

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zo.infolist():
                data = overrides.get(item.filename)
                if data is None:
                    data = zo.read(item.filename)
                elif isinstance(data, str):
                    data = data.encode("utf-8")
                zout.writestr(item.filename, data)
            for name, data in additions.items():
                zout.writestr(name, data)
        result = buf.getvalue()

        # Validasi: zip utuh & masih bisa dibuka openpyxl
        if zipfile.ZipFile(io.BytesIO(result)).testzip() is not None:
            return output_bytes, "hasil pemulihan tidak lolos validasi zip"
        openpyxl.load_workbook(io.BytesIO(result), read_only=True).close()
        return result, "OK"
    except Exception as e:
        return output_bytes, f"{type(e).__name__}: {e}"


if st.button("🚀 Process & Generate Template", type="primary"):
    if not csv_file or not prev_so_file:
        st.error("Dokslinya diperhatikan lagi lohya 💀")
    else:
        with st.spinner("Suki tidur besok SO lagi sama member 😹"):
            try:
                template_bytes = fetch_master_template()
                template_raw = template_bytes.getvalue()
                wb = openpyxl.load_workbook(io.BytesIO(template_raw))
            except Exception as e:
                st.error(
                    f"❌ Gagal mengambil Template Master dari Google Drive: {e}"
                )
                st.stop()

            df_raw = load_raw_inventory_data(csv_file)
            prev_so_map = build_prev_so_dict(prev_so_file)

            # Ekstraksi LOC Code aktif dari Summary
            active_summary_locs = [
                sdata["loc_code"] for sdata in summary_data_inputs
            ]

            # Filter Batch Missing
            missing_batch_set = extract_missing_batch_set(
                missing_file, target_loc_list=active_summary_locs
            )

            # Hapus Sheet1 & Sheet2 jika ada
            for sname in ["Sheet1", "Sheet2", "sheet1", "sheet2"]:
                if sname in wb.sheetnames:
                    del wb[sname]

            # Filter Hapus Baris Missing
            initial_count = len(df_raw)
            if missing_batch_set:

                def is_not_missing(row):
                    raw_batch = get_val(
                        row, "BATCH", "BATCH NO", "BATCH_NO", default=""
                    )
                    norm_batch = normalize_batch(raw_batch)
                    return norm_batch not in missing_batch_set

                df_raw = df_raw[
                    df_raw.apply(is_not_missing, axis=1)
                ].reset_index(drop=True)

            deleted_missing_count = initial_count - len(df_raw)

            # Filter NLA vs Regular
            nla_condition = pd.Series([False] * len(df_raw), index=df_raw.index)
            for col_check in [
                "STATUS",
                "REMARK",
                "REASON",
                "CATEGORY",
                "CAT",
                "CONDITION",
            ]:
                if col_check in df_raw.columns:
                    nla_condition = nla_condition | (
                        df_raw[col_check]
                        .astype(str)
                        .str.upper()
                        .str.contains("NLA")
                    )

            df_nla = df_raw[nla_condition].reset_index(drop=True)
            df_worksheet = df_raw[~nla_condition].reset_index(drop=True)

            # -----------------------------------------------------
            # UPDATE SHEET WORKSHEET
            # -----------------------------------------------------
            if "Worksheet" in wb.sheetnames:
                ws = wb["Worksheet"]
                set_cell_safe(ws, 3, 3, f" {station_input}")
                set_cell_safe(ws, 4, 3, f" {location_input}")
                set_cell_safe(ws, 5, 3, f" {periode_input}")
                write_so_table_to_sheet(ws, df_worksheet, prev_so_map)

            # -----------------------------------------------------
            # UPDATE SHEET SUMMARY
            # -----------------------------------------------------
            if "Summary" in wb.sheetnames:
                ws_sum = wb["Summary"]
                set_cell_safe(ws_sum, 3, 4, f": {station_input}")  # D3: Station
                set_cell_safe(ws_sum, 4, 4, f": {periode_input}")  # D4: Periode

                start_sum_row = 11
                for i, sdata in enumerate(summary_data_inputs):
                    r_curr = start_sum_row + i
                    set_cell_safe(ws_sum, r_curr, 2, i + 1)  # B: NO
                    set_cell_safe(
                        ws_sum, r_curr, 3, sdata["division"]
                    )  # C: DIVISION
                    set_cell_safe(ws_sum, r_curr, 4, sdata["pic"])  # D: PIC
                    set_cell_safe(
                        ws_sum, r_curr, 5, sdata["loc_code"]
                    )  # E: LOC CODE
                    set_cell_safe(
                        ws_sum, r_curr, 6, sdata["loc_desc"]
                    )  # F: LOCATION DESCRIPTION

            # PivotTable di-refresh otomatis saat file dibuka di Excel (opsional)
            if REFRESH_PIVOT_ON_OPEN:
                refresh_pivots_on_open(wb)
            wb.calculation.fullCalcOnLoad = True

            # Save ke memory buffer untuk download
            output_buffer = io.BytesIO()
            wb.save(output_buffer)

            # Kembalikan logo header cetak yang dibuang openpyxl
            if RESTORE_HEADER_LOGO:
                logo_bytes, logo_status = restore_header_logo(
                    template_raw, output_buffer.getvalue()
                )
            else:
                logo_bytes, logo_status = output_buffer.getvalue(), "DIMATIKAN"
            output_buffer = io.BytesIO(logo_bytes)
            output_buffer.seek(0)

            st.success("Worksheet Berhasil Dihitamkan 😹")
            if logo_status == "OK":
                st.caption("🖼️ Logo header berhasil dipulihkan.")
            elif logo_status == "DIMATIKAN":
                st.caption("🖼️ Pemulihan logo header sedang dimatikan (RESTORE_HEADER_LOGO = False).")
            else:
                st.warning(f"🖼️ Logo header tidak bisa dipulihkan: {logo_status}")
            if deleted_missing_count > 0:
                st.info(
                    f"🗑️ Sebanyak **{deleted_missing_count}** baris barang berstatus **MISSING** terdampak Ragnamok 😹🙏"
                )

            st.download_button(
                label="📥 Download Hasil Excel Stock Opname",
                data=output_buffer,
                file_name=f"Worksheet_Stock_Opname_{station_input}_Updated.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
