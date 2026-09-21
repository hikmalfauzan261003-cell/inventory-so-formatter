import streamlit as st
import pandas as pd
import openpyxl
import io

st.set_page_config(
    page_title="Stock Opname Worksheet Formatter",
    page_icon="📦",
    layout="wide"
)

st.title("📦 Excel Formatter - Worksheet Template SO_2")
st.write("Upload file template (`TEMPLATE SO_2.xlsx`) dan file Data Utama (`REPORT INVENTORY...`) untuk memperbarui sheet `worksheet` tanpa mengubah baris 1-7.")

# Sidebar Upload File
st.sidebar.header("📁 Upload File")
template_file = st.sidebar.file_uploader("1. Upload Template (`TEMPLATE SO_2.xlsx`)", type=["xlsx"])
data_file = st.sidebar.file_uploader("2. Upload Data Utama (Excel/CSV)", type=["xlsx", "xls", "csv"])

if template_file is not None and data_file is not None:
    try:
        # 1. Baca Data Utama
        if data_file.name.endswith(".csv"):
            df_main = pd.read_csv(data_file, sep=None, engine="python")
        else:
            df_main = pd.read_excel(data_file)
            
        # Standarisasi nama kolom data utama
        df_main.columns = [str(c).strip().upper() for c in df_main.columns]

        # 2. Urutan Kolom Target di Sheet 'worksheet'
        target_columns = [
            "LOCATION", "BIN", "GRB", "BATCH", "PN", "SN", 
            "PN DESCRIPTION", "QTY AVAILABLE", "QTY RESERVED", "QTY IN TRANSFER", 
            "QTY PENDING RI", "QTY US", "QTY IN REPAIR", "SHELF LIFE EXPIRATION", 
            "TOOL LIFE EXPIRATION", "CONDITION", "CATEGORY", "OWNER", "UOM"
        ]

        # Aturan Mapping Kolom
        mapping = {
            "LOCATION": ["LOCATION", "LOC", "STORE"],
            "BIN": ["BIN", "LOCATION BIN"],
            "GRB": ["GRB"],
            "BATCH": ["BATCH", "BATCH NO", "BATCH_NO"],
            "PN": ["PN", "PART NO", "PART NUMBER", "P/N"],
            "SN": ["SN", "SERIAL NO", "SERIAL NUMBER", "S/N"],
            "PN DESCRIPTION": ["PN DESCRIPTION", "DESCRIPTION", "PART DESCRIPTION"],
            "QTY AVAILABLE": ["QTY AVAILABLE", "QTY EMRO", "QTY", "QUANTITY"],
            "QTY RESERVED": ["QTY RESERVED", "RESERVED"],
            "QTY IN TRANSFER": ["QTY IN TRANSFER", "QTY IN TF", "TRANSFER"],
            "QTY PENDING RI": ["QTY PENDING RI", "PENDING RI"],
            "QTY US": ["QTY US", "UNSERVICEABLE"],
            "QTY IN REPAIR": ["QTY IN REPAIR", "REPAIR"],
            "SHELF LIFE EXPIRATION": ["SHELF LIFE EXPIRATION", "SHELF LIFE EXP", "EXPIRY DATE"],
            "TOOL LIFE EXPIRATION": ["TOOL LIFE EXPIRATION", "TOOL LIFE EXP"],
            "CONDITION": ["CONDITION", "COND"],
            "CATEGORY": ["CATEGORY", "CATEGORY GROUP"],
            "OWNER": ["OWNER"],
            "UOM": ["UOM", "UNIT"]
        }

        # Petakan Data Utama ke Kolom Target
        df_mapped = pd.DataFrame()
        for col in target_columns:
            sources = mapping.get(col, [col])
            matched_col = next((s for s in sources if s in df_main.columns), None)
            if matched_col:
                df_mapped[col] = df_main[matched_col]
            else:
                df_mapped[col] = ""

        # Isi nilai 0 untuk kolom numerik jika kosong
        numeric_cols = ["QTY AVAILABLE", "QTY RESERVED", "QTY IN TRANSFER", "QTY PENDING RI", "QTY US", "QTY IN REPAIR"]
        for nc in numeric_cols:
            df_mapped[nc] = pd.to_numeric(df_mapped[nc], errors='coerce').fillna(0)

        st.subheader("📋 Preview Data yang Akan Dimasukkan ke Sheet `worksheet`")
        st.dataframe(df_mapped.head(10), use_container_width=True)

        # 3. Masukkan Data ke File Template Asli menggunakan openpyxl (Preserve Baris 1-7 & Sheet lain)
        wb = openpyxl.load_workbook(template_file)
        
        if "worksheet" in wb.sheetnames:
            ws = wb["worksheet"]
            
            # Hapus data lama dari baris ke-8 ke bawah (jika ada)
            if ws.max_row >= 8:
                ws.delete_rows(8, amount=ws.max_row - 7)

            # Tulis data baru mulai dari baris ke-8
            start_row = 8
            for r_idx, row in df_mapped.iterrows():
                for c_idx, val in enumerate(row, start=1):
                    ws.cell(row=start_row + r_idx, column=c_idx, value=val)
            
            # Simpan file ke memory buffer
            output_buffer = io.BytesIO()
            wb.save(output_buffer)
            output_buffer.seek(0)

            st.success("✅ Sheet `worksheet` berhasil diperbarui! Baris 1-7 & sheet lainnya tetap utuh.")

            # Download Button
            st.download_button(
                label="📥 Download Template SO_2 (Updated)",
                data=output_buffer,
                file_name="TEMPLATE_SO_2_UPDATED.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        else:
            st.error("Sheet bernama `worksheet` tidak ditemukan di dalam file template yang diupload.")

    except Exception as e:
        st.error(f"Terjadi kesalahan saat memproses file: {e}")